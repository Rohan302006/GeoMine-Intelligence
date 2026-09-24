import os
import io
from datetime import datetime
from typing import List, Dict, Any, Optional
import pandas as pd
from reportlab.lib.pagesizes import letter, A4
from reportlab.lib import colors
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, HRFlowable
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

from app.core.config import settings

class ReportGenerator:
    def __init__(self):
        self.output_dir = settings.REPORTS_DIR
        os.makedirs(self.output_dir, exist_ok=True)

    def generate_pdf_report(self, title: str, report_type: str, summary_text: str, table_data: List[List[str]], sources: List[str]) -> str:
        filename = f"Report_{datetime.now().strftime('%Y%m%d_%H%M%S')}.pdf"
        file_path = os.path.join(self.output_dir, filename)
        
        doc = SimpleDocTemplate(file_path, pagesize=A4, rightMargin=36, leftMargin=36, topMargin=36, bottomMargin=36)
        elements = []
        styles = getSampleStyleSheet()
        
        # Institutional Header Style
        title_style = ParagraphStyle(
            'GovHeader',
            parent=styles['Heading1'],
            fontName='Helvetica-Bold',
            fontSize=16,
            leading=20,
            textColor=colors.HexColor('#0F2E59'),
            alignment=1
        )
        
        sub_style = ParagraphStyle(
            'GovSub',
            parent=styles['Normal'],
            fontName='Helvetica',
            fontSize=10,
            leading=14,
            textColor=colors.HexColor('#4B5563'),
            alignment=1
        )
        
        body_style = ParagraphStyle(
            'GovBody',
            parent=styles['Normal'],
            fontName='Helvetica',
            fontSize=10,
            leading=15,
            textColor=colors.HexColor('#1F2937')
        )
        
        elements.append(Paragraph("GEOMINE INTELLIGENCE DIRECTORATE / MINERAL ANALYTICS DIVISION", title_style))
        elements.append(Paragraph("CENTRAL MINE PLANNING & DESIGN INSTITUTE LIMITED (CMPDI) / COAL INDIA LIMITED", sub_style))
        elements.append(Spacer(1, 10))
        elements.append(HRFlowable(width="100%", thickness=1.5, color=colors.HexColor('#0F2E59'), spaceAfter=15))
        
        elements.append(Paragraph(f"<b>REPORT TITLE:</b> {title.upper()}", styles['Heading2']))
        elements.append(Paragraph(f"<b>Report Category:</b> {report_type} | <b>Generated On:</b> {datetime.now().strftime('%d-%b-%Y %H:%M')}", sub_style))
        elements.append(Spacer(1, 12))
        
        elements.append(Paragraph("<b>1. EXECUTIVE SUMMARY & AI SYNTHESIS</b>", styles['Heading3']))
        elements.append(Paragraph(summary_text, body_style))
        elements.append(Spacer(1, 15))
        
        if table_data:
            elements.append(Paragraph("<b>2. STATISTICAL PERFORMANCE ANNEXURE</b>", styles['Heading3']))
            table = Table(table_data)
            table.setStyle(TableStyle([
                ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#0F2E59')),
                ('TEXTCOLOR', (0, 0), (-1, 0), colors.whitesmoke),
                ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
                ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
                ('FONTSIZE', (0, 0), (-1, -1), 9),
                ('BOTTOMPADDING', (0, 0), (-1, 0), 8),
                ('TOPPADDING', (0, 0), (-1, -1), 6),
                ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#D1D5DB')),
                ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.HexColor('#F9FAFB'), colors.white])
            ]))
            elements.append(table)
            elements.append(Spacer(1, 15))
            
        elements.append(Paragraph("<b>3. SOURCE CITATIONS & AUDIT REPOSITORY</b>", styles['Heading3']))
        for src in sources:
            elements.append(Paragraph(f"• {src}", sub_style))
        elements.append(Spacer(1, 15))
        
        disclaimer = ParagraphStyle(
            'Disclaimer',
            parent=styles['Italic'],
            fontSize=8,
            leading=10,
            textColor=colors.HexColor('#6B7280'),
            alignment=1
        )
        elements.append(HRFlowable(width="100%", thickness=0.5, color=colors.HexColor('#9CA3AF'), spaceAfter=8))
        elements.append(Paragraph("This platform is a prototype solution developed for the Smart India Hackathon 2026 (PS ID: SIH26023). It is not an official portal of the Ministry of Coal or CMPDI/CIL.", disclaimer))
        
        doc.build(elements)
        return file_path

    def generate_docx_report(self, title: str, report_type: str, summary_text: str, table_data: List[List[str]], sources: List[str]) -> str:
        filename = f"Report_{datetime.now().strftime('%Y%m%d_%H%M%S')}.docx"
        file_path = os.path.join(self.output_dir, filename)
        
        doc = docx.Document()
        
        h1 = doc.add_heading("GEOMINE INTELLIGENCE DIRECTORATE — MINERAL ANALYTICS DIVISION", level=1)
        h1.alignment = WD_ALIGN_PARAGRAPH.CENTER
        
        p = doc.add_paragraph("CENTRAL MINE PLANNING & DESIGN INSTITUTE LIMITED / COAL INDIA LIMITED")
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        
        doc.add_heading(title, level=2)
        doc.add_paragraph(f"Category: {report_type} | Date: {datetime.now().strftime('%d-%b-%Y %H:%M')}")
        
        doc.add_heading("Executive Summary", level=3)
        doc.add_paragraph(summary_text)
        
        if table_data:
            doc.add_heading("Statistical Data Table", level=3)
            table = doc.add_table(rows=len(table_data), cols=len(table_data[0]))
            table.style = 'Table Grid'
            for r_idx, row in enumerate(table_data):
                for c_idx, val in enumerate(row):
                    table.cell(r_idx, c_idx).text = str(val)
                    
        doc.add_heading("Data Sources & Verification", level=3)
        for s in sources:
            doc.add_paragraph(f"• {s}")
            
        doc.add_paragraph("\nDisclaimer: This platform is a prototype solution developed for the Smart India Hackathon 2026 (PS ID: SIH26023). It is not an official portal of the Ministry of Coal or CMPDI/CIL.")
        doc.save(file_path)
        return file_path

    def generate_excel_report(self, title: str, table_data: List[List[str]]) -> str:
        filename = f"Report_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
        file_path = os.path.join(self.output_dir, filename)
        
        if table_data and len(table_data) > 1:
            headers = table_data[0]
            rows = table_data[1:]
            df = pd.DataFrame(rows, columns=headers)
        else:
            df = pd.DataFrame({"Metric": ["No data available"]})
            
        with pd.ExcelWriter(file_path, engine='openpyxl') as writer:
            df.to_excel(writer, sheet_name="Data", index=False)
            
        return file_path

report_generator = ReportGenerator()
